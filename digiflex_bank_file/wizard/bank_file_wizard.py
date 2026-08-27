from odoo import models, fields

# -*- coding: utf-8 -*-
from odoo import http
from odoo.http import request
import io
import json
import xlsxwriter
import base64
import openpyxl
import os
import shutil
from odoo.exceptions import UserError

class BankFileDownloadWizard(models.TransientModel):
    _name = 'bank.file.wizard'
    _description = 'Digiflex Bank File Download Wizard'

    payment_ids = fields.Many2many(
        'account.payment',
        string='Payments',
        default=lambda self: self.env.context.get('payment_ids', []),
        readonly=True,
    )
    journal_id = fields.Many2one(
        'account.journal',
        string='Journal',
        default=lambda self: self.env.context.get('journal_id', False),
        readonly=True,
    )

    output_format = fields.Selection(
        [
            ('default', ' '), 
            ('bnc.xlsm', 'BNC - Banco Nacional de Crédito'), 
            ('bdv.xlsm', 'BDV - Banco de Venezuela')
        ],
        string='Formato de Salida',
        default='default',
        required=True
    )

    attachment_id = fields.Many2one('ir.attachment', 
        'XLSX Attachment', 
        readonly=True)

    def generate_bnc(self, sheet, payments):

        for idx, rec in enumerate(payments, start=3):

            company_bank_account = False
            if rec.journal_id.bank_account_id.acc_number:
                company_bank_account = ''.join(ch for ch in rec.journal_id.bank_account_id.acc_number if ch.isdigit())

            partner_bank_account = False
            if rec.partner_bank_id.acc_number:
                partner_bank_account = ''.join(ch for ch in rec.partner_bank_id.acc_number if ch.isdigit())
            full_rif = rec.partner_id.vat or rec.partner_id.rif or ''
            rif = full_rif.replace('-', '')
            ref = rec.ref.replace('/', ' ').replace('-', ' ') 

            sheet.cell(row=idx, column=1, value=rec.date.strftime('%d/%m/%Y') if rec.date else '')  # Fecha de Pago
            sheet.cell(row=idx, column=2, value=company_bank_account or '')  # Cuenta a Debitar
            sheet.cell(row=idx, column=3, value=partner_bank_account or '')  # Cuenta Beneficiario
            sheet.cell(row=idx, column=4, value=rec.amount or '')  # Monto
            sheet.cell(row=idx, column=5, value=ref or '')  # Descripción
            sheet.cell(row=idx, column=6, value=rif) # ID Beneficiario
            sheet.cell(row=idx, column=7, value=rec.partner_id.name or '')  # Nombre Beneficiario
            sheet.cell(row=idx, column=8, value=rec.partner_id.email or '')  # Email Beneficiario
            sheet.cell(row=idx, column=9, value='') # Referencia del Cliente

    def generate_bdv(self, sheet, payments):
        for idx, rec in enumerate(payments, start=11):

            full_rif = rec.partner_id.vat.replace('-', '') if rec.partner_id.vat else ''
            rif = full_rif[1:] if full_rif else ''
            payment_number = ''.join(ch for ch in (rec.name or '') if ch.isdigit())
            payment_number = payment_number[-8:] if len(payment_number) > 8 else payment_number
            acc_number = rec.partner_bank_id.acc_number.replace('-', '') if rec.partner_bank_id.acc_number else ''

            sheet.cell(row=idx, column=2, value=rec.partner_id.name) # 2.- NOMBRE DEL BENEFICIARIO
            sheet.cell(row=idx, column=3, value=payment_number) # 3.- N° DE REFERENCIA DEL CRÉDITO
            sheet.cell(row=idx, column=4, value=rec.partner_id.vat[0] if rec.partner_id.vat else '')  # 4.- LETRA RIF/CI
            sheet.cell(row=idx, column=5, value=rif or '')  # 5.- NÚMERO RIF/CI
            sheet.cell(row=idx, column=6, value='C')# 6.- TIPO DE CUENTA
            sheet.cell(row=idx, column=7, value=acc_number)  # 7.- N° DE CUENTA DEL BENEFICIARIO
            sheet.cell(row=idx, column=8, value=rec.amount or '') # 8.- MONTO DEL CRÉDITO
            sheet.cell(row=idx, column=9, value=2) # 9.- TIPO DE PAGO
            sheet.cell(row=idx, column=10, value=rec.partner_bank_id.bank_id.bank_code_vzla or 36) # 10.- BANCO
            sheet.cell(row=idx, column=11, value='') # 11.- DURACIÓN DEL CHEQUE
            sheet.cell(row=idx, column=12, value=str(rec.partner_id.email or '') )  # 12.- EMAIL DEL BENEFICIARIO
            sheet.cell(row=idx, column=13, value=rec.date.strftime('%d/%m/%Y') if rec.date else '') # 13.- FECHA VALOR DEL DÉBITO



    def action_download_file(self):
        if self.output_format not in ('bnc.xlsm', 'bdv.xlsm'):
            raise UserError("Selecciona un formato de archivo válido")

        module_path = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
        template_path = os.path.join(module_path, 'data', f'{self.output_format}')
        name, file_extension = os.path.splitext(os.path.basename(template_path))
        file_name = f"{name.upper()}_{fields.Date.context_today(self).strftime('%Y-%m-%d')}{file_extension}"

        dest_path = os.path.join('/tmp', file_name)
        if os.path.exists(dest_path):
            os.remove(dest_path)
        shutil.copyfile(template_path, dest_path)
        file_path = dest_path
        wb = openpyxl.load_workbook(file_path, keep_vba=True)
        if len(wb.worksheets) < 2:
            return
       
        payments = request.env['account.payment'].browse(self.payment_ids.ids)

        if self.output_format == 'bnc.xlsm':
            sheet = wb.worksheets[1]
            self.generate_bnc(sheet, payments)
        elif self.output_format == 'bdv.xlsm':
            sheet = wb.worksheets[2]
            self.generate_bdv(sheet, payments)
       
        wb.save(file_path)
        with open(file_path, 'rb') as f:
            data = f.read()
        b64data = base64.b64encode(data).decode()
        new_attachment = self.env['ir.attachment'].create({
            'name': file_path.split('/')[-1],
            'type': 'binary',
            'datas': b64data,
            'mimetype': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        })
        return {
            'type': 'ir.actions.act_url',
            'url': f'/web/content/{new_attachment.id}?download=true'
        }