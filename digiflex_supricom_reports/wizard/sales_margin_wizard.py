# -*- coding: utf-8 -*-
import base64
import io
from datetime import date
from dateutil.relativedelta import relativedelta
from odoo import api, fields, models, _
from odoo.exceptions import UserError

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False


class DigiflexSalesMarginWizard(models.TransientModel):
    _name = 'digiflex.sales.margin.wizard'
    _description = 'Asistente - Reporte de Ventas con Utilidad y Margen'

    date_from = fields.Date(
        string='Del',
        required=True,
        default=lambda self: date.today().replace(day=1),
    )
    date_to = fields.Date(
        string='Al',
        required=True,
        default=fields.Date.context_today,
    )
    sale_condition = fields.Selection([
        ('all', 'Ambos'),
        ('cash', 'Contado'),
        ('credit', 'Crédito'),
    ], string='Condición', default='all', required=True)

    movement_type = fields.Selection([
        ('all', 'Ambos'),
        ('sales', 'Ventas'),
        ('refunds', 'Devoluciones'),
    ], string='Tipo Movimiento', default='all', required=True)

    sale_type = fields.Selection([
        ('all', 'Ambos'),
        ('products', 'Productos'),
        ('services', 'Servicios'),
    ], string='Tipo de Venta', default='all', required=True)

    include_service_cost = fields.Boolean(
        string='Costo de Servicios',
        default=False,
        help='Indica si se deben considerar costos en los ítems clasificados como servicios.',
    )

    company_id = fields.Many2one(
        'res.company',
        string='Compañía',
        default=lambda self: self.env.company,
        required=True,
    )

    def _get_report_domain(self):
        self.ensure_one()
        domain = [
            ('company_id', '=', self.company_id.id),
            ('invoice_date', '>=', self.date_from),
            ('invoice_date', '<=', self.date_to),
        ]

        if self.sale_condition != 'all':
            domain.append(('sale_condition', '=', self.sale_condition))

        if self.movement_type != 'all':
            domain.append(('movement_type', '=', self.movement_type))

        if self.sale_type == 'products':
            domain.append(('product_type', '=', 'product'))
        elif self.sale_type == 'services':
            domain.append(('product_type', '=', 'service'))

        return domain

    def action_view_report(self):
        self.ensure_one()
        domain = self._get_report_domain()

        return {
            'name': _('Ventas / Devoluciones - Margen y Utilidad'),
            'type': 'ir.actions.act_window',
            'res_model': 'digiflex.sales.margin.report',
            'view_mode': 'tree,pivot,graph',
            'domain': domain,
            'context': {
                'search_default_group_by_date': 1,
                'search_default_group_by_salesperson': 0,
            },
            'target': 'current',
        }

    def action_export_excel(self):
        self.ensure_one()
        if not HAS_OPENPYXL:
            raise UserError(_("La librería 'openpyxl' no está instalada en el servidor."))

        domain = self._get_report_domain()
        records = self.env['digiflex.sales.margin.report'].search(domain, order='invoice_date asc, document_number asc')

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Ventas y Devoluciones"
        ws.views.sheetView[0].showGridLines = True

        # Styles
        title_font = Font(name="Calibri", size=14, bold=True, color="1B5E20")
        header_fill = PatternFill(start_color="C8E6C9", end_color="C8E6C9", fill_type="solid")
        filter_fill = PatternFill(start_color="E8F5E9", end_color="E8F5E9", fill_type="solid")
        th_fill = PatternFill(start_color="A5D6A7", end_color="A5D6A7", fill_type="solid")
        total_fill = PatternFill(start_color="C8E6C9", end_color="C8E6C9", fill_type="solid")

        bold_font = Font(name="Calibri", size=10, bold=True)
        th_font = Font(name="Calibri", size=10, bold=True, color="1B5E20")
        data_font = Font(name="Calibri", size=10)

        thin_side = Side(style="thin", color="CCCCCC")
        border_all = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
        double_bottom = Border(top=thin_side, bottom=Side(style="double", color="1B5E20"))

        align_left = Alignment(horizontal="left", vertical="center")
        align_center = Alignment(horizontal="center", vertical="center")
        align_right = Alignment(horizontal="right", vertical="center")

        # 1. Main Title
        ws.merge_cells("A1:J1")
        ws["A1"] = "VENTAS / DEVOLUCIONES - REPORTE DE MARGEN SOBRE COSTO"
        ws["A1"].font = title_font
        ws["A1"].alignment = align_left

        # 2. Filter summary bar (matching Image 2 header)
        ws.merge_cells("A3:B3")
        ws["A3"] = "Rango de Fechas:"
        ws["A3"].font = bold_font
        ws["C3"] = f"Del: {self.date_from.strftime('%d/%m/%Y')}  Al: {self.date_to.strftime('%d/%m/%Y')}"
        ws["C3"].font = data_font

        ws["E3"] = "Condición:"
        ws["E3"].font = bold_font
        ws["F3"] = dict(self._fields['sale_condition'].selection).get(self.sale_condition, 'Ambos')
        ws["F3"].font = data_font

        ws["H3"] = "Tipo Movimiento:"
        ws["H3"].font = bold_font
        ws["I3"] = dict(self._fields['movement_type'].selection).get(self.movement_type, 'Ambos')
        ws["I3"].font = data_font

        ws.merge_cells("A4:B4")
        ws["A4"] = "Tipo de Venta:"
        ws["A4"].font = bold_font
        ws["C4"] = dict(self._fields['sale_type'].selection).get(self.sale_type, 'Ambos')
        ws["C4"].font = data_font

        ws["E4"] = "Costo de Servicios:"
        ws["E4"].font = bold_font
        ws["F4"] = "Sí" if self.include_service_cost else "No"
        ws["F4"].font = data_font

        for r in range(3, 5):
            for c in range(1, 11):
                ws.cell(row=r, column=c).fill = filter_fill

        # 3. Table Headers
        headers = [
            "Fecha", "Numero", "Cuenta", "Cliente", "Vendedor",
            "Unidades", "Costo", "Venta", "Utilidad", "MargenC (%)"
        ]

        row_idx = 6
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=header)
            cell.font = th_font
            cell.fill = th_fill
            cell.alignment = align_center
            cell.border = border_all

        # 4. Data Rows
        start_data_row = 7
        row_idx = start_data_row

        for rec in records:
            # Handle cost for services if disabled
            cost = rec.cost_total
            if rec.is_service and not self.include_service_cost:
                cost = 0.0

            sale = rec.sale_total
            profit = sale - cost
            margin_pct = (profit / cost * 100.0) if cost != 0 else (100.0 if sale > 0 else 0.0)

            ws.cell(row=row_idx, column=1, value=rec.invoice_date.strftime('%Y/%m/%d') if rec.invoice_date else '').alignment = align_center
            ws.cell(row=row_idx, column=2, value=rec.document_number or '/').alignment = align_left
            ws.cell(row=row_idx, column=3, value=rec.partner_ref or '').alignment = align_center
            ws.cell(row=row_idx, column=4, value=rec.partner_name or '').alignment = align_left
            ws.cell(row=row_idx, column=5, value=rec.salesperson_name or '').alignment = align_left

            c_units = ws.cell(row=row_idx, column=6, value=rec.quantity)
            c_units.number_format = '#,##0.00'
            c_units.alignment = align_right

            c_cost = ws.cell(row=row_idx, column=7, value=cost)
            c_cost.number_format = '#,##0.00'
            c_cost.alignment = align_right

            c_sale = ws.cell(row=row_idx, column=8, value=sale)
            c_sale.number_format = '#,##0.00'
            c_sale.alignment = align_right

            c_profit = ws.cell(row=row_idx, column=9, value=profit)
            c_profit.number_format = '#,##0.00'
            c_profit.alignment = align_right

            c_margin = ws.cell(row=row_idx, column=10, value=margin_pct / 100.0)
            c_margin.number_format = '0.00%'
            c_margin.alignment = align_right

            for col_idx in range(1, 11):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.font = data_font
                cell.border = border_all

            row_idx += 1

        # 5. Totals Row
        end_data_row = row_idx - 1
        if end_data_row >= start_data_row:
            ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=5)
            total_label = ws.cell(row=row_idx, column=1, value="TOTALES")
            total_label.font = bold_font
            total_label.alignment = align_right

            t_units = ws.cell(row=row_idx, column=6, value=f"=SUM(F{start_data_row}:F{end_data_row})")
            t_units.number_format = '#,##0.00'
            t_units.font = bold_font
            t_units.alignment = align_right

            t_cost = ws.cell(row=row_idx, column=7, value=f"=SUM(G{start_data_row}:G{end_data_row})")
            t_cost.number_format = '#,##0.00'
            t_cost.font = bold_font
            t_cost.alignment = align_right

            t_sale = ws.cell(row=row_idx, column=8, value=f"=SUM(H{start_data_row}:H{end_data_row})")
            t_sale.number_format = '#,##0.00'
            t_sale.font = bold_font
            t_sale.alignment = align_right

            t_profit = ws.cell(row=row_idx, column=9, value=f"=SUM(I{start_data_row}:I{end_data_row})")
            t_profit.number_format = '#,##0.00'
            t_profit.font = bold_font
            t_profit.alignment = align_right

            t_margin = ws.cell(row=row_idx, column=10, value=f"=IF(G{row_idx}>0, I{row_idx}/G{row_idx}, 0)")
            t_margin.number_format = '0.00%'
            t_margin.font = bold_font
            t_margin.alignment = align_right

            for col_idx in range(1, 11):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.fill = total_fill
                cell.border = double_bottom

        # Auto-adjust column widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        ws.column_dimensions['A'].width = 14
        ws.column_dimensions['B'].width = 20
        ws.column_dimensions['C'].width = 12
        ws.column_dimensions['D'].width = 30
        ws.column_dimensions['E'].width = 22
        ws.column_dimensions['F'].width = 14
        ws.column_dimensions['G'].width = 16
        ws.column_dimensions['H'].width = 16
        ws.column_dimensions['I'].width = 16
        ws.column_dimensions['J'].width = 14

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        file_content = output.getvalue()
        output.close()

        file_name = f"Ventas_Margen_{self.date_from.strftime('%Y%m%d')}_{self.date_to.strftime('%Y%m%d')}.xlsx"
        attachment = self.env['ir.attachment'].create({
            'name': file_name,
            'type': 'binary',
            'datas': base64.b64encode(file_content),
            'mimetype': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            'res_model': self._name,
            'res_id': self.id,
        })

        return {
            'type': 'ir.actions.act_url',
            'url': f'/web/content/{attachment.id}?download=true',
            'target': 'self',
        }
