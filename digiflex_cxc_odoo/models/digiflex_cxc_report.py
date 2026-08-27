# -*- coding: utf-8 -*-
import base64
import io
import logging
import re
from odoo import models, fields, tools, _
from odoo.exceptions import UserError

_logger = logging.getLogger(__name__)

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

_ILLEGAL_CHARACTERS_RE = re.compile(r'[\x00-\x08\x0B-\x0C\x0E-\x1F]')


def clean_str(val):
    if not val:
        return ""
    if not isinstance(val, str):
        val = str(val)
    return _ILLEGAL_CHARACTERS_RE.sub('', val)


class DigiflexCxcReport(models.Model):
    _name = 'digiflex.cxc.report'
    _description = 'Reporte de Cuentas por Cobrar'
    _auto = False
    _order = 'days_overdue desc, partner_name asc'

    move_id = fields.Many2one('account.move', string='Factura / Asiento', readonly=True)
    move_line_id = fields.Many2one('account.move.line', string='Línea Contable', readonly=True)
    partner_id = fields.Many2one('res.partner', string='Cliente ID', readonly=True)
    user_id = fields.Many2one('res.users', string='Vendedor ID', readonly=True)
    company_id = fields.Many2one('res.company', string='Compañía ID', readonly=True)
    currency_id = fields.Many2one('res.currency', string='Moneda ID', readonly=True)

    partner_name = fields.Char(string='Cliente', readonly=True)
    user_name = fields.Char(string='Vendedor', readonly=True)
    company_name = fields.Char(string='Compañía', readonly=True)
    phone = fields.Char(string='Teléfono', readonly=True)

    transaction_type = fields.Char(string='Transacción', readonly=True)
    document_number = fields.Char(string='Documento', readonly=True)
    nro_ctrl = fields.Char(string='Número de Control', readonly=True)

    invoice_date = fields.Date(string='Fecha', readonly=True)
    date_maturity = fields.Date(string='Fecha Vencimiento', readonly=True)
    days_overdue = fields.Integer(string='Días Vencidos', readonly=True, group_operator=False)

    amount_residual = fields.Monetary(string='Saldo', currency_field='currency_id', readonly=True)
    amount_current = fields.Monetary(string='Por Vencer', currency_field='currency_id', readonly=True)
    amount_1_30 = fields.Monetary(string='1-30', currency_field='currency_id', readonly=True)
    amount_31_60 = fields.Monetary(string='31-60', currency_field='currency_id', readonly=True)
    amount_61_90 = fields.Monetary(string='61-90', currency_field='currency_id', readonly=True)
    amount_91_plus = fields.Monetary(string='91+', currency_field='currency_id', readonly=True)

    def init(self):
        tools.drop_view_if_exists(self.env.cr, self._table)
        has_fel_pa = tools.column_exists(self.env.cr, 'account_move', 'fel_pa_document_number')
        doc_num_expr = "COALESCE(NULLIF(am.fel_pa_document_number, ''), am.name, am.ref, '/')" if has_fel_pa else "COALESCE(am.name, am.ref, '/')"

        self.env.cr.execute("""
            CREATE OR REPLACE VIEW %s AS (
                SELECT
                    aml.id AS id,
                    aml.id AS move_line_id,
                    am.id AS move_id,
                    am.partner_id AS partner_id,
                    COALESCE(rp.name, '/') AS partner_name,
                    COALESCE(rp.phone, rp.mobile) AS phone,
                    COALESCE(am.invoice_user_id, rp.user_id) AS user_id,
                    COALESCE(rp_user.name, ru.login, '/') AS user_name,
                    aml.company_id AS company_id,
                    COALESCE(rc.name, '/') AS company_name,
                    aml.company_currency_id AS currency_id,
                    CASE 
                        WHEN am.move_type = 'out_invoice' THEN 'Factura'
                        WHEN am.move_type = 'out_refund' THEN 'Nota de Crédito'
                        WHEN am.move_type = 'out_receipt' THEN 'Recibo'
                        WHEN am.move_type = 'entry' THEN 'Asiento'
                        ELSE 'Otro'
                    END AS transaction_type,
                    %s AS document_number,
                    am.nro_ctrl AS nro_ctrl,
                    COALESCE(am.invoice_date, aml.date) AS invoice_date,
                    COALESCE(aml.date_maturity, aml.date, am.invoice_date) AS date_maturity,
                    (CURRENT_DATE - COALESCE(aml.date_maturity, aml.date, am.invoice_date)) AS days_overdue,
                    aml.amount_residual AS amount_residual,
                    CASE 
                        WHEN (CURRENT_DATE - COALESCE(aml.date_maturity, aml.date, am.invoice_date)) <= 0 
                        THEN aml.amount_residual ELSE 0.0 
                    END AS amount_current,
                    CASE 
                        WHEN (CURRENT_DATE - COALESCE(aml.date_maturity, aml.date, am.invoice_date)) BETWEEN 1 AND 30 
                        THEN aml.amount_residual ELSE 0.0 
                    END AS amount_1_30,
                    CASE 
                        WHEN (CURRENT_DATE - COALESCE(aml.date_maturity, aml.date, am.invoice_date)) BETWEEN 31 AND 60 
                        THEN aml.amount_residual ELSE 0.0 
                    END AS amount_31_60,
                    CASE 
                        WHEN (CURRENT_DATE - COALESCE(aml.date_maturity, aml.date, am.invoice_date)) BETWEEN 61 AND 90 
                        THEN aml.amount_residual ELSE 0.0 
                    END AS amount_61_90,
                    CASE 
                        WHEN (CURRENT_DATE - COALESCE(aml.date_maturity, aml.date, am.invoice_date)) > 90 
                        THEN aml.amount_residual ELSE 0.0 
                    END AS amount_91_plus
                FROM account_move_line aml
                JOIN account_move am ON am.id = aml.move_id
                JOIN account_account aa ON aa.id = aml.account_id
                LEFT JOIN res_partner rp ON rp.id = am.partner_id
                LEFT JOIN res_users ru ON ru.id = COALESCE(am.invoice_user_id, rp.user_id)
                LEFT JOIN res_partner rp_user ON rp_user.id = ru.partner_id
                LEFT JOIN res_company rc ON rc.id = aml.company_id
                WHERE aa.account_type = 'asset_receivable'
                  AND am.state = 'posted'
                  AND (aml.reconciled IS NOT TRUE OR aml.reconciled IS NULL)
                  AND aml.amount_residual != 0
            )
        """ % (self._table, doc_num_expr))

    def action_export_excel(self):
        if not HAS_OPENPYXL:
            raise UserError(_("La librería 'openpyxl' no está disponible en el servidor."))

        _logger.info("[DIGIFLEX CXC] action_export_excel iniciado. Contexto: %s", self._context)
        raw_domain = self._context.get('active_domain', [])
        _logger.info("[DIGIFLEX CXC] active_domain recibido: %s (tipo: %s)", raw_domain, type(raw_domain))

        if isinstance(raw_domain, str):
            try:
                import ast
                raw_domain = ast.literal_eval(raw_domain)
            except Exception as e:
                _logger.warning("[DIGIFLEX CXC] Error al evaluar raw_domain: %s", e)
                raw_domain = []

        domain = list(raw_domain) if isinstance(raw_domain, (list, tuple)) else []
        domain.append(('company_id', 'in', self.env.companies.ids))
        _logger.info("[DIGIFLEX CXC] Dominio final de búsqueda: %s", domain)

        records = self.env['digiflex.cxc.report'].sudo().search(domain)
        _logger.info("[DIGIFLEX CXC] Registros obtenidos para exportar: %s", len(records))

        # 2. Crear libro de trabajo Excel
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Cuentas por Cobrar"
        ws.views.sheetView[0].showGridLines = True

        # Estilos profesionales
        font_title = Font(name="Calibri", size=14, bold=True, color="1F4E78")
        font_subtitle = Font(name="Calibri", size=10, italic=True, color="595959")
        font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        font_data = Font(name="Calibri", size=10)
        font_total = Font(name="Calibri", size=11, bold=True, color="000000")

        fill_header = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        fill_total = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")

        align_center = Alignment(horizontal="center", vertical="center")
        align_left = Alignment(horizontal="left", vertical="center")
        align_right = Alignment(horizontal="right", vertical="center")

        border_thin = Side(border_style="thin", color="D9D9D9")
        border_top_thin = Side(border_style="thin", color="000000")
        border_bottom_double = Side(border_style="double", color="000000")

        cell_border = Border(left=border_thin, right=border_thin, top=border_thin, bottom=border_thin)
        total_border = Border(top=border_top_thin, bottom=border_bottom_double)

        # Título y metadatos
        company_name = clean_str(self.env.company.name)
        today_date = fields.Date.context_today(self)
        today_str = today_date.strftime("%d/%m/%Y")

        ws.cell(row=1, column=1, value=f"REPORTE DE CUENTAS POR COBRAR - {company_name.upper()}").font = font_title
        ws.cell(row=2, column=1, value=f"Fecha de emisión: {today_str} | Registros: {len(records)}").font = font_subtitle

        # Encabezados de tabla
        headers = [
            "Vendedor", "Cliente", "Teléfono", "Transacción", "Documento",
            "Número de Control", "Fecha Emisión", "Fecha Vencimiento",
            "Días Vencidos", "Saldo Total", "Por Vencer",
            "1-30 Días", "31-60 Días", "61-90 Días", "+90 Días"
        ]

        for col_num, header_title in enumerate(headers, 1):
            cell = ws.cell(row=4, column=col_num, value=header_title)
            cell.font = font_header
            cell.fill = fill_header
            cell.alignment = align_center

        ws.row_dimensions[4].height = 25

        # Llenar datos de cada fila utilizando campos directos de la vista SQL
        current_row = 5
        for rec in records:
            ws.cell(row=current_row, column=1, value=clean_str(rec.user_name)).alignment = align_left
            ws.cell(row=current_row, column=2, value=clean_str(rec.partner_name)).alignment = align_left
            ws.cell(row=current_row, column=3, value=clean_str(rec.phone)).alignment = align_center
            ws.cell(row=current_row, column=4, value=clean_str(rec.transaction_type)).alignment = align_center
            ws.cell(row=current_row, column=5, value=clean_str(rec.document_number)).alignment = align_left
            ws.cell(row=current_row, column=6, value=clean_str(rec.nro_ctrl)).alignment = align_center
            
            c_date = ws.cell(row=current_row, column=7, value=rec.invoice_date.strftime("%d/%m/%Y") if rec.invoice_date else "")
            c_date.alignment = align_center
            
            c_mat = ws.cell(row=current_row, column=8, value=rec.date_maturity.strftime("%d/%m/%Y") if rec.date_maturity else "")
            c_mat.alignment = align_center
            
            c_days = ws.cell(row=current_row, column=9, value=rec.days_overdue or 0)
            c_days.alignment = align_center
            c_days.number_format = "0"

            monetary_cols = [
                (10, rec.amount_residual),
                (11, rec.amount_current),
                (12, rec.amount_1_30),
                (13, rec.amount_31_60),
                (14, rec.amount_61_90),
                (15, rec.amount_91_plus)
            ]
            for c_idx, val in monetary_cols:
                cell_m = ws.cell(row=current_row, column=c_idx, value=val or 0.0)
                cell_m.alignment = align_right
                cell_m.number_format = "#,##0.00"

            for col_num in range(1, 16):
                c = ws.cell(row=current_row, column=col_num)
                c.font = font_data
                c.border = cell_border

            current_row += 1

        # Totales Generales
        ws.cell(row=current_row, column=1, value="TOTAL GENERAL").font = font_total
        ws.cell(row=current_row, column=1).alignment = align_left

        for col_idx in range(10, 16):
            col_letter = get_column_letter(col_idx)
            cell_t = ws.cell(row=current_row, column=col_idx, value=f"=SUM({col_letter}5:{col_letter}{current_row-1})")
            cell_t.font = font_total
            cell_t.alignment = align_right
            cell_t.number_format = "#,##0.00"

        for col_idx in range(1, 16):
            c = ws.cell(row=current_row, column=col_idx)
            c.fill = fill_total
            c.border = total_border

        ws.row_dimensions[current_row].height = 22

        # Autoajustar ancho de columnas
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = clean_str(str(cell.value or ''))
                if cell.number_format == "#,##0.00":
                    val_str += "   "
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        # Generar binario BytesIO
        output = io.BytesIO()
        wb.save(output)
        excel_bytes = output.getvalue()
        output.close()

        # Adjunto temporal creado con sudo()
        filename = f"CxC_{today_date.strftime('%Y-%m-%d')}.xlsx"
        attachment = self.env['ir.attachment'].sudo().create({
            'name': filename,
            'datas': base64.b64encode(excel_bytes),
            'mimetype': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        })

        return {
            'type': 'ir.actions.act_url',
            'url': f'/web/content/{attachment.id}?download=true',
            'target': 'self',
        }
